"""Replace master slides 74-88 (Productivity) with the rebuilt Productivity deck (ADP headcount), 2026-10-20."""
import copy, shutil
from pptx import Presentation
import os
D=os.path.join(os.path.dirname(os.path.abspath(__file__)),"..","decks")+"/"
MASTER=D+"Fixed Ops - Master Review Deck.pptx"
shutil.copy(MASTER, D+"Fixed Ops - Master Review Deck (backup before swap).pptx")
prs=Presentation(MASTER); src=Presentation(D+"Fixed Ops Productivity - House Style.pptx")
assert len(src.slides)==15 and len(prs.slides)==88
layout=prs.slide_layouts[0]; P="{http://schemas.openxmlformats.org/presentationml/2006/main}"
RID="{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"; pkg=prs.part.package; moved=set()
def clone(ss):
    s=prs.slides.add_slide(layout); dst=s._element.cSld
    bg=ss._element.cSld.find(P+"bg")
    if bg is not None: dst.insert(0,copy.deepcopy(bg))
    tree=dst.find(P+"spTree")
    for ch in list(tree): tree.remove(ch)
    for ch in ss._element.cSld.find(P+"spTree"): tree.append(copy.deepcopy(ch))
    for el in tree.iter():
        rid=el.get(RID)
        if rid is None: continue
        rel=ss.part.rels[rid]; part=rel.target_part
        if part not in moved:
            part.partname=pkg.next_partname("/ppt/charts/chart%d.xml")
            for r2 in part.rels.values():
                if not r2.is_external and "embeddings" in str(r2.target_part.partname):
                    ext=str(r2.target_part.partname).rsplit(".",1)[-1]
                    r2.target_part.partname=pkg.next_partname("/ppt/embeddings/Microsoft_Excel_Worksheet%d."+ext)
            moved.add(part)
        el.set(RID,s.part.relate_to(part,rel.reltype))
    if ss.has_notes_slide and ss.notes_slide.notes_text_frame.text.strip():
        s.notes_slide.notes_text_frame.text=ss.notes_slide.notes_text_frame.text
    return s
new=[clone(ss) for ss in src.slides]
lst=prs.slides._sldIdLst; els=list(lst)
old=els[73:88]; newels=els[88:]
for el in old+newels: lst.remove(el)
for el in newels: lst.insert(73+newels.index(el), el) if False else None
# rebuild order: first 73, new 15
keep=[e for e in list(lst)]
for e in list(lst): lst.remove(e)
for e in keep[:73]+newels: lst.append(e)
for e in old: prs.part.drop_rel(e.rId)
prs.save(MASTER); print("master slides:",len(Presentation(MASTER).slides))
