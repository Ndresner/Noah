#!/usr/bin/env python3
"""Findlay ADP Workforce Now CLI.

Pulls ADP data over the ADP API and writes .xlsx files in the same layout as the manual
ADP report exports, so the existing report scripts take them without changes:

  roster      Active roster + headcount by store/job title (Roster, Headcount sheets)
  earnings    "Earnings - Parts and Service Employees" layout (YTD Gross Pay) -> pay-vs-gross build_base.py --adp
  tech-hours  "Tech Employee List / Tech Efficiency Report Data" layout -> proficiency build_data.py
  probe       Checks credentials and which ADP APIs this account can call
  seed-config Builds the job-title filters from last period's manual exports
  compare     Diffs an API-built file against a manual export of the same report/period

Credentials come from environment variables only (never files in this repo):
  ADP_CLIENT_ID, ADP_CLIENT_SECRET, ADP_CERT, ADP_KEY
ADP_CERT / ADP_KEY are either file paths or the PEM text itself.
"""
import argparse
import datetime as dt
import json
import os
import re
import sys
import tempfile
import time

API_BASE = "https://api.adp.com"
TOKEN_URL = "https://accounts.adp.com/auth/oauth/v2/token"
HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_CONFIG = os.path.join(HERE, "adp_config.json")
EXAMPLE_CONFIG = os.path.join(HERE, "adp_config.example.json")

# Column layouts of the manual ADP exports. Order matters: the proficiency loader unpacks by position.
EARNINGS_COLS = ["Company Code", "File Number", "Payroll Name", "Job Title Description", "Gross Pay",
                 "Position Status", "Hire/Rehire Date"]
TECH_COLS = ["Location Description", "Position ID", "Legal Last Name", "Legal First Name",
             "Job Title Description", "Position Status", "Regular Hours Total", "Overtime Hours Total"]
ROSTER_COLS = ["Company Code", "File Number", "Position ID", "Payroll Name", "Legal Last Name", "Legal First Name",
               "Job Title Description", "Position Status", "Location Description", "Department",
               "Hire/Rehire Date", "Original Hire Date", "Termination Date", "Associate OID"]
STATUS_NAMES = {"A": "Active", "L": "Leave", "T": "Terminated", "I": "Inactive", "D": "Deceased", "R": "Retired"}


class ADPError(RuntimeError):
    pass


# ---------------------------------------------------------------- helpers
def g(obj, *path, default=None):
    """Safe nested get: g(w, "person", "legalName", "givenName"); ints index lists."""
    for p in path:
        if isinstance(p, int):
            if not isinstance(obj, list) or len(obj) <= p:
                return default
            obj = obj[p]
        else:
            if not isinstance(obj, dict) or p not in obj:
                return default
            obj = obj[p]
    return default if obj is None else obj


def code_text(c):
    """ADP code objects carry codeValue / shortName / longName; return the most readable one."""
    if isinstance(c, str):
        return c
    return g(c, "longName") or g(c, "shortName") or g(c, "codeValue") or ""


def parse_date(s):
    if not s:
        return None
    if isinstance(s, (dt.date, dt.datetime)):
        return s if isinstance(s, dt.date) and not isinstance(s, dt.datetime) else s.date()
    return dt.date.fromisoformat(str(s)[:10])


def amount(x):
    v = g(x, "amountValue") if isinstance(x, dict) else x
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _pem(value, suffix):
    """Env value is a path or PEM text. PEM text goes to a private temp file for requests."""
    if not value:
        return None
    if "-----BEGIN" not in value:
        if not os.path.exists(value):
            raise ADPError(f"{suffix} file not found: {value}")
        return value
    fd, path = tempfile.mkstemp(suffix=suffix)
    os.chmod(path, 0o600)
    with os.fdopen(fd, "w") as fh:
        fh.write(value.replace("\\n", "\n"))
    return path


# ---------------------------------------------------------------- client
class ADPClient:
    def __init__(self, client_id, client_secret, cert=None, key=None, session=None,
                 api_base=API_BASE, token_url=TOKEN_URL, pause=0.0, log=None):
        if session is None:
            import requests
            session = requests.Session()
        self.s = session
        if cert:
            self.s.cert = (cert, key) if key else cert
        self.cid, self.secret = client_id, client_secret
        self.api_base, self.token_url = api_base.rstrip("/"), token_url
        self.pause = pause
        self.log = log or (lambda m: print(m, file=sys.stderr))
        self._tok, self._exp = None, 0.0
        self.calls = 0

    @classmethod
    def from_env(cls, **kw):
        cid, sec = os.environ.get("ADP_CLIENT_ID"), os.environ.get("ADP_CLIENT_SECRET")
        missing = [n for n, v in (("ADP_CLIENT_ID", cid), ("ADP_CLIENT_SECRET", sec),
                                  ("ADP_CERT", os.environ.get("ADP_CERT"))) if not v]
        if missing:
            raise ADPError(f"missing environment variables: {', '.join(missing)} "
                           "(set them as environment secrets, never in the repo)")
        return cls(cid, sec, _pem(os.environ.get("ADP_CERT"), ".pem"), _pem(os.environ.get("ADP_KEY"), ".key"), **kw)

    def token(self, force=False):
        if self._tok and not force and time.time() < self._exp - 60:
            return self._tok
        r = self.s.post(self.token_url, data={"grant_type": "client_credentials",
                                              "client_id": self.cid, "client_secret": self.secret}, timeout=60)
        if r.status_code != 200:
            # Never echo the request; the body only carries ADP's error code/description.
            raise ADPError(f"token request failed: HTTP {r.status_code} {r.text[:300]}")
        j = r.json()
        self._tok, self._exp = j["access_token"], time.time() + float(j.get("expires_in", 3600))
        return self._tok

    def get(self, path, params=None):
        """GET with one token refresh on 401 and backoff on 429/5xx. 204 / empty body -> None."""
        url = path if path.startswith("http") else self.api_base + path
        refreshed = False
        for attempt in range(5):
            if self.pause:
                time.sleep(self.pause)
            self.calls += 1
            r = self.s.get(url, params=params, timeout=120,
                           headers={"Authorization": f"Bearer {self.token()}", "Accept": "application/json"})
            if r.status_code == 401 and not refreshed:
                self.token(force=True); refreshed = True; continue
            if r.status_code in (429, 500, 502, 503, 504):
                wait = float(r.headers.get("Retry-After") or 2 ** (attempt + 1))
                self.log(f"  HTTP {r.status_code} on {path}; retrying in {wait:.0f}s")
                time.sleep(wait); continue
            if r.status_code == 204 or (r.status_code == 200 and not r.content):
                return None
            if r.status_code != 200:
                raise ADPError(f"GET {path} failed: HTTP {r.status_code} {r.text[:300]}")
            return r.json()
        raise ADPError(f"GET {path} failed after retries")

    def paged(self, path, key, top=100, params=None):
        skip, out = 0, []
        while True:
            p = dict(params or {}, **{"$top": top, "$skip": skip})
            j = self.get(path, p)
            batch = (j or {}).get(key) or []
            out.extend(batch)
            if len(batch) < top:
                return out
            skip += top

    # --- endpoints
    def workers(self):
        return self.paged("/hr/v2/workers", "workers")

    def pay_statements(self, aoid, last_n):
        j = self.get(f"/payroll/v1/workers/{aoid}/pay-statements", {"numberoflastpaydates": last_n})
        return (j or {}).get("payStatements") or []

    def pay_statement_detail(self, stmt):
        href = g(stmt, "payDetailUri", "href") or g(stmt, "statementImageUri", "href")
        if not href or "pay-statements/" not in href:
            return stmt
        j = self.get(href)
        return (j or {}).get("payStatement") or stmt


# ---------------------------------------------------------------- worker parsing
def primary_assignment(w):
    was = w.get("workAssignments") or []
    for wa in was:
        if wa.get("primaryIndicator"):
            return wa
    return was[0] if was else {}


def org_unit(wa, type_name):
    for u in (wa.get("homeOrganizationalUnits") or []) + (wa.get("assignedOrganizationalUnits") or []):
        if code_text(u.get("typeCode")).strip().lower() == type_name.lower():
            return code_text(u.get("nameCode"))
    return ""


def worker_row(w):
    wa = primary_assignment(w)
    name = g(w, "person", "legalName") or {}
    last, first = name.get("familyName1") or "", name.get("givenName") or ""
    mid = name.get("middleName") or ""
    formatted = name.get("formattedName") or (f"{last}, {first}" + (f" {mid[:1]}" if mid else "")).strip(", ")
    status = g(wa, "assignmentStatus", "statusCode") or g(w, "workerStatus", "statusCode") or {}
    status_txt = status.get("longName") or status.get("shortName") or STATUS_NAMES.get(status.get("codeValue"), status.get("codeValue") or "")
    title = wa.get("jobTitle") or code_text(wa.get("jobCode"))
    orig_hire = parse_date(g(w, "workerDates", "originalHireDate") or wa.get("hireDate"))
    rehire = parse_date(g(w, "workerDates", "rehireDate"))
    return {
        "aoid": w.get("associateOID", ""),
        "company_code": (wa.get("payrollGroupCode") or "").strip(),
        "file_number": (wa.get("payrollFileNumber") or "").strip(),
        "position_id": (wa.get("positionID") or "").strip(),
        "payroll_name": formatted,
        "last": last, "first": first,
        "title": (title or "").strip(),
        "status": status_txt,
        "location": code_text(g(wa, "homeWorkLocation", "nameCode")).strip(),
        "department": org_unit(wa, "Department"),
        "hire": rehire if rehire and (not orig_hire or rehire > orig_hire) else orig_hire,
        "orig_hire": orig_hire,
        "term": parse_date(g(w, "workerDates", "terminationDate") or wa.get("terminationDate")),
    }


def keep_worker(r, flt, active_only=False, period_start=None):
    """Job-title / company-code filter. Terminated workers stay if they could have pay in the period."""
    titles = {t.strip().upper() for t in flt.get("job_titles", [])}
    codes = {c.strip().upper() for c in flt.get("company_codes", [])}
    excl = {t.strip().upper() for t in flt.get("exclude_job_titles", [])}
    t = r["title"].upper()
    if titles and t not in titles:
        return False
    if t in excl:
        return False
    if codes and r["company_code"].upper() not in codes:
        return False
    if active_only and r["status"].lower() not in ("active", "leave"):
        return False
    if period_start and r["term"] and r["term"] < period_start:
        return False
    return True


# ---------------------------------------------------------------- pay statements
def classify_earning(e, hours_cfg):
    code = (code_text(g(e, "earningCode")) + " " + str(g(e, "earningCode", "codeValue", default=""))).upper()
    name = str(e.get("earningCodeName") or "").upper()
    blob = f"{code} {name}"

    def hit(keys):  # whole-word match, so "OT" doesn't catch "SPOT BONUS"
        return any(re.search(r"(?<![A-Z0-9])" + re.escape(k.upper()) + r"(?![A-Z0-9])", blob) for k in keys)
    if hit(hours_cfg.get("overtime_codes", [])):
        return "ot"
    if hit(hours_cfg.get("regular_codes", [])):
        return "reg"
    return "other"


def statement_hours(detail, hours_cfg):
    reg = ot = 0.0
    other = {}
    for e in detail.get("earnings") or []:
        h = float(g(e, "payHours", "hoursQuantity", default=0) or 0)
        if not h:
            continue
        k = classify_earning(e, hours_cfg)
        if k == "reg":
            reg += h
        elif k == "ot":
            ot += h
        else:
            nm = e.get("earningCodeName") or code_text(g(e, "earningCode")) or "?"
            other[nm] = other.get(nm, 0) + h
    return reg, ot, other


def in_range(d, start, end):
    return d is not None and start <= d <= end


def stmt_date(s, basis):
    if basis == "period_end":
        return parse_date(g(s, "payPeriod", "endDate")) or parse_date(s.get("payDate"))
    return parse_date(s.get("payDate"))


def lookback(start, today=None):
    """Pay dates to request: enough weekly pay dates to reach back past `start`."""
    today = today or dt.date.today()
    return max(4, min(104, (today - start).days // 7 + 3))


# ---------------------------------------------------------------- xlsx output
def write_xlsx(path, sheets):
    """sheets: list of (name, header, rows). Plain data extract, no formulas."""
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, header, rows in sheets:
        ws = wb.create_sheet(name)
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append(r)
        for i, h in enumerate(header, 1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if r[i - 1] is not None] or [8])
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = min(width + 3, 45)
            if rows and isinstance(rows[0][i - 1], dt.date):
                for c in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                    for cell in c:
                        cell.number_format = "m/d/yyyy"
            elif rows and isinstance(rows[0][i - 1], float):
                for c in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                    for cell in c:
                        cell.number_format = "#,##0.00"
        ws.freeze_panes = "A2"
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)


def require_titles(cfg, key):
    flt = cfg.get(key, {})
    if not flt.get("job_titles"):
        raise ADPError(f"config '{key}.job_titles' is empty - run seed-config with last period's manual export first "
                       "(an empty filter would pull every employee in the group)")
    return flt


def load_config(path):
    p = path or (DEFAULT_CONFIG if os.path.exists(DEFAULT_CONFIG) else EXAMPLE_CONFIG)
    with open(p) as fh:
        return json.load(fh)


def dump_raw(args, name, obj):
    if args.dump_raw:
        os.makedirs(args.dump_raw, exist_ok=True)
        with open(os.path.join(args.dump_raw, name), "w") as fh:
            json.dump(obj, fh, indent=1, default=str)


# ---------------------------------------------------------------- commands
def fetch_workers(client, args):
    ws = client.workers()
    dump_raw(args, "workers.json", ws)
    rows = [worker_row(w) for w in ws]
    client.log(f"workers: {len(rows)} from ADP")
    return rows


def cmd_roster(args, client, cfg):
    rows = fetch_workers(client, args)
    flt = cfg.get("roster", {})
    rows = [r for r in rows if keep_worker(r, flt, active_only=not args.include_terminated)]
    rows.sort(key=lambda r: (r["company_code"], r["payroll_name"]))
    data = [[r["company_code"], r["file_number"], r["position_id"], r["payroll_name"], r["last"], r["first"],
             r["title"], r["status"], r["location"], r["department"], r["hire"], r["orig_hire"], r["term"], r["aoid"]]
            for r in rows]
    hc = {}
    for r in rows:
        if r["status"].lower() in ("active", "leave"):
            k = (r["company_code"], r["location"], r["title"])
            hc[k] = hc.get(k, 0) + 1
    hc_rows = [[k[0], k[1], k[2], v] for k, v in sorted(hc.items())]
    write_xlsx(args.out, [("Roster", ROSTER_COLS, data),
                          ("Headcount", ["Company Code", "Location Description", "Job Title Description", "Headcount"], hc_rows)])
    client.log(f"wrote {args.out}: {len(data)} employees, {sum(hc.values())} active headcount")


def collect_pay(client, args, cfg, workers, start, end, basis, need_detail):
    """Per worker: pay statements whose date (pay date or period end) falls in [start, end]."""
    n = args.lookback or lookback(start)
    out, raw, errors = {}, {}, []
    for i, r in enumerate(workers, 1):
        if i % 50 == 0:
            client.log(f"  pay statements: {i}/{len(workers)}")
        try:
            stmts = [s for s in client.pay_statements(r["aoid"], n) if in_range(stmt_date(s, basis), start, end)]
            if need_detail or basis == "period_end":
                stmts = [client.pay_statement_detail(s) for s in stmts]
                stmts = [s for s in stmts if in_range(stmt_date(s, basis), start, end)]
        except ADPError as e:
            errors.append(f"{r['payroll_name']} ({r['aoid']}): {e}")
            continue
        out[r["aoid"]] = stmts
        raw[r["aoid"]] = stmts
    dump_raw(args, "pay_statements.json", raw)
    if errors:
        client.log(f"WARNING: {len(errors)} workers failed; first: {errors[0]}")
        if len(errors) > max(5, len(workers) // 20):
            raise ADPError("too many pay-statement failures (>5%) - output not written")
    return out, errors


def cmd_earnings(args, client, cfg):
    # The manual export's Gross Pay is year-to-date through the pay check date (build_base.py divides it by
    # months employed since Jan 1), so default to summing every pay statement dated Jan 1..pay date.
    end = parse_date(args.pay_date)
    start = parse_date(args.pay_date_from) if args.pay_date_from else dt.date(end.year, 1, 1)
    flt = require_titles(cfg, "earnings")
    workers = [r for r in fetch_workers(client, args) if keep_worker(r, flt, period_start=start - dt.timedelta(days=31))]
    client.log(f"earnings: {len(workers)} workers match the job-title filter; pay dates {start}..{end}")
    pay, errors = collect_pay(client, args, cfg, workers, start, end, "pay_date", need_detail=False)
    rows = []
    for r in workers:
        stmts = pay.get(r["aoid"], [])
        if not stmts and not args.include_zero:
            continue
        gross = round(sum(amount(s.get("grossPayAmount")) for s in stmts), 2)
        rows.append([r["company_code"], r["file_number"], r["payroll_name"], r["title"], gross, r["status"], r["hire"]])
    rows.sort(key=lambda x: (x[0], x[2]))
    write_xlsx(args.out, [("1", EARNINGS_COLS, rows)])
    client.log(f"wrote {args.out}: {len(rows)} employees, gross pay ${sum(x[4] for x in rows):,.2f}; "
               f"{len(errors)} errors; {client.calls} API calls")


def cmd_tech_hours(args, client, cfg):
    start, end = parse_date(args.start), parse_date(args.end)
    flt = require_titles(cfg, "tech_hours")
    hcfg = cfg.get("hours", {})
    workers = [r for r in fetch_workers(client, args) if keep_worker(r, flt, period_start=start)]
    client.log(f"tech-hours: {len(workers)} workers match; {args.basis} in {start}..{end}")
    pay, errors = collect_pay(client, args, cfg, workers, start, end, args.basis, need_detail=True)
    rows, other_all = [], {}
    for r in workers:
        stmts = pay.get(r["aoid"], [])
        if not stmts and not args.include_zero:
            continue
        reg = ot = 0.0
        for s in stmts:
            a, b, other = statement_hours(s, hcfg)
            reg += a; ot += b
            for k, v in other.items():
                other_all[k] = other_all.get(k, 0) + v
        rows.append([r["location"], r["position_id"], r["last"], r["first"], r["title"], r["status"],
                     round(reg, 2), round(ot, 2)])
    rows.sort(key=lambda x: (x[0], x[2], x[3]))
    write_xlsx(args.out, [("1", TECH_COLS, rows)])
    client.log(f"wrote {args.out}: {len(rows)} employees, {sum(x[6] for x in rows):,.1f} reg + "
               f"{sum(x[7] for x in rows):,.1f} OT hrs; {len(errors)} errors; {client.calls} API calls")
    if other_all:
        client.log("hours NOT counted as Regular/Overtime (check hours.regular_codes / overtime_codes in config):")
        for k, v in sorted(other_all.items(), key=lambda kv: -kv[1]):
            client.log(f"  {k}: {v:,.1f}")


def cmd_probe(args, client, cfg):
    checks = []
    try:
        client.token()
        checks.append(("OAuth token (accounts.adp.com)", "OK"))
    except Exception as e:
        print(f"OAuth token: FAILED - {e}")
        return 1
    aoid = None
    try:
        j = client.get("/hr/v2/workers", {"$top": 1})
        ws = (j or {}).get("workers") or []
        aoid = ws[0].get("associateOID") if ws else None
        checks.append(("Workers API /hr/v2/workers", "OK" if ws else "OK (no rows)"))
        if ws and args.show_sample:
            print(json.dumps(worker_row(ws[0]), default=str, indent=1))
    except Exception as e:
        checks.append(("Workers API /hr/v2/workers", f"FAILED - {e}"))
    if aoid:
        try:
            st = client.pay_statements(aoid, 1)
            checks.append(("Pay statements /payroll/v1/.../pay-statements", "OK" if st else "OK (none for sample worker)"))
            if st:
                d = client.pay_statement_detail(st[0])
                checks.append(("Pay statement detail (earnings + hours)", "OK" if d.get("earnings") else "no earnings lines"))
        except Exception as e:
            checks.append(("Pay statements /payroll/v1/.../pay-statements", f"FAILED - {e}"))
    for name, res in checks:
        print(f"{name:<52} {res}")
    return 0 if all(r.startswith("OK") for _, r in checks) else 1


def read_sheet(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["1"] if "1" in wb.sheetnames else wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h is not None else "" for h in next(it)]
    rows = [dict(zip(hdr, r)) for r in it]
    return hdr, rows


def cmd_seed_config(args, client, cfg):
    out = load_config(args.config)
    for key, path in (("earnings", args.earnings_export), ("tech_hours", args.tech_export)):
        if not path:
            continue
        _, rows = read_sheet(path)
        titles = sorted({str(r.get("Job Title Description")).strip() for r in rows if r.get("Job Title Description")})
        out.setdefault(key, {})["job_titles"] = titles
        print(f"{key}: {len(titles)} job titles from {os.path.basename(path)}")
    target = args.config or DEFAULT_CONFIG
    with open(target, "w") as fh:
        json.dump(out, fh, indent=2)
    print(f"wrote {target}")


def cmd_compare(args, client, cfg):
    hdr_a, a = read_sheet(args.api)
    hdr_b, b = read_sheet(args.export)
    keycols = [c for c in ("Position ID",) if c in hdr_a and c in hdr_b] or \
              [c for c in ("Company Code", "File Number") if c in hdr_a and c in hdr_b] or \
              [c for c in ("Company Code", "Payroll Name") if c in hdr_a and c in hdr_b]
    if not keycols:
        print("no shared key columns"); return 1

    def key(r):
        return tuple(str(r.get(c) or "").strip().upper() for c in keycols)
    ia = {key(r): r for r in a if any(key(r))}
    ib = {key(r): r for r in b if any(key(r)) and (r.get("Job Title Description") or "Position ID" in keycols)}
    only_a, only_b = sorted(set(ia) - set(ib)), sorted(set(ib) - set(ia))
    cols = [c for c in hdr_b if c in hdr_a and c not in keycols]
    diffs = {c: [] for c in cols}
    for k in sorted(set(ia) & set(ib)):
        for c in cols:
            va, vb = ia[k].get(c), ib[k].get(c)
            if isinstance(va, (int, float)) or isinstance(vb, (int, float)):
                try:
                    if abs(float(va or 0) - float(vb or 0)) <= args.tolerance:
                        continue
                except (TypeError, ValueError):
                    pass
            elif parse_cell(va) == parse_cell(vb):
                continue
            diffs[c].append((k, va, vb))
    print(f"key: {' + '.join(keycols)} | API rows {len(ia)} | export rows {len(ib)} | matched {len(set(ia) & set(ib))}")
    print(f"only in API file: {len(only_a)}   only in manual export: {len(only_b)}")
    for k in only_b[:15]:
        print(f"  missing from API: {k} {ib[k].get('Payroll Name') or ib[k].get('Legal Last Name')}")
    for k in only_a[:15]:
        print(f"  extra in API:     {k} {ia[k].get('Payroll Name') or ia[k].get('Legal Last Name')}")
    for c in cols:
        if c in ("Gross Pay", "Regular Hours Total", "Overtime Hours Total"):
            ta = sum(float(r.get(c) or 0) for r in ia.values()); tb = sum(float(r.get(c) or 0) for r in ib.values())
            print(f"{c}: API total {ta:,.2f} vs export {tb:,.2f} (diff {ta - tb:,.2f})")
        print(f"{c}: {len(diffs[c])} mismatched rows")
        for k, va, vb in diffs[c][:5]:
            print(f"    {k}: API={va!r} export={vb!r}")
    clean = not only_a and not only_b and not any(diffs.values())
    print("RESULT:", "MATCH" if clean else "DIFFERENCES - review before using the API file")
    return 0 if clean else 2


def parse_cell(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return parse_date(v)
    s = str(v or "").strip().upper()
    try:
        return dt.datetime.strptime(s, "%m/%d/%Y").date()
    except ValueError:
        return s


def build_parser():
    ap = argparse.ArgumentParser(prog="adp", description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--config", help="filters/earning-code config JSON (default adp_config.json, else the example)")
    ap.add_argument("--dump-raw", metavar="DIR", help="also save the raw ADP JSON here (contains pay data - scratchpad only)")
    ap.add_argument("--pause", type=float, default=0.0, help="seconds between API calls (rate limiting)")
    sp = ap.add_subparsers(dest="cmd", required=True)

    p = sp.add_parser("probe", help="check credentials and API access")
    p.add_argument("--show-sample", action="store_true", help="print one parsed worker (contains PII)")

    p = sp.add_parser("roster", help="roster + headcount")
    p.add_argument("--out", required=True)
    p.add_argument("--include-terminated", action="store_true")

    p = sp.add_parser("earnings", help="Earnings - Parts and Service Employees layout")
    p.add_argument("--pay-date", required=True, help="pay check date YYYY-MM-DD (Gross Pay is YTD through this date)")
    p.add_argument("--pay-date-from", help="first pay date to include (default Jan 1 of the pay-date year)")
    p.add_argument("--out", required=True)
    p.add_argument("--include-zero", action="store_true", help="keep matched workers with no pay in range")
    p.add_argument("--lookback", type=int, help="pay dates to request per worker (default: computed)")

    p = sp.add_parser("tech-hours", help="Tech Efficiency Report Data layout (Reg + OT hours)")
    p.add_argument("--start", required=True); p.add_argument("--end", required=True)
    p.add_argument("--basis", choices=["period_end", "pay_date"], default="period_end",
                   help="which date must fall in start..end (default: pay period end date)")
    p.add_argument("--out", required=True)
    p.add_argument("--include-zero", action="store_true")
    p.add_argument("--lookback", type=int)

    p = sp.add_parser("seed-config", help="set job-title filters from last period's manual exports")
    p.add_argument("--earnings-export"); p.add_argument("--tech-export")

    p = sp.add_parser("compare", help="diff an API-built file against a manual export")
    p.add_argument("--api", required=True); p.add_argument("--export", required=True)
    p.add_argument("--tolerance", type=float, default=0.01)
    return ap


def main(argv=None):
    args = build_parser().parse_args(argv)
    cfg = load_config(args.config)
    offline = args.cmd in ("seed-config", "compare")
    try:
        client = None if offline else ADPClient.from_env(pause=args.pause)
        fn = {"probe": cmd_probe, "roster": cmd_roster, "earnings": cmd_earnings, "tech-hours": cmd_tech_hours,
              "seed-config": cmd_seed_config, "compare": cmd_compare}[args.cmd]
        return fn(args, client, cfg) or 0
    except ADPError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
