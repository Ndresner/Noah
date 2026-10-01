"""Build the base Pay-vs-Group-Average workbook from a raw ADP export.

Implements Earnings_Parts_Service_Pay_Report_BUILD_SPEC (the original Aug-2026 build):
Master Summary, Store Ranker, Raw Data, Store & Position Summary, Group Position Summary,
and one tab per store. All figures are live formulas off Raw Data.

The output is the input (--base) for add_gross_workload.py. Run recalc.py on the final file.
"""
import argparse, datetime as dt
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.formatting.rule import Rule, ColorScaleRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.hyperlink import Hyperlink

ap = argparse.ArgumentParser()
ap.add_argument("--adp", required=True, help="ADP 'Earnings - Parts and Service Employees' export (.xlsx)")
ap.add_argument("--out", required=True)
ap.add_argument("--period-start", default=None, help="YYYY-MM-DD, default Jan 1 of period-end year")
ap.add_argument("--period-end", required=True, help="YYYY-MM-DD; hire cutoff and months-employed end")
ap.add_argument("--exclude", default="EXPRESS SERVICE MANAGER,LEAD SALES PORTER")
ap.add_argument("--qlik", default=None, help="optional Qlik JSON; its advisor_internal + advisor_overrides move confirmed "
                "internal (recon) Service Advisors to INTERNAL SERVICE ADVISOR")
A = ap.parse_args()
PE = dt.date.fromisoformat(A.period_end)
PS = dt.date.fromisoformat(A.period_start) if A.period_start else dt.date(PE.year, 1, 1)
EXCL = {x.strip().upper() for x in A.exclude.split(",") if x.strip()}
PE_TXT = f"{PE.month}/{PE.day}/{PE.year % 100:02d}"
DPE = f"DATE({PE.year},{PE.month},{PE.day})"; DPS = f"DATE({PS.year},{PS.month},{PS.day})"

# ADP company code -> store name, exactly as in the ADP lookup (typos and "Lincoln " trailing space kept).
# Used when the export has no 'Sheet1' lookup tab; a 'Sheet1' tab, if present, overrides/extends it.
CODE_MAP = {
    "H4X": "Acura", "HUZ": "Audi Henderson", "D4K": "Audi Reno Tahoe", "HLD": "CJDR Post Falls", "YD6": "Cadillac",
    "D7A": "Chevrolet Las Vegas", "J01": "Customs", "GMW": "GMC Prescott", "HHU": "Honda Flagstaff",
    "H4V": "Honda Henderson", "YD9": "Honda North", "A0E": "Honda Spokane", "L64": "Hyundai Sg. George",
    "I78": "Hyundal Prescott", "M7T": "INEOS Grenadier", "ZC7": "Jaguar Land Rover Las Vegas", "ZFJ": "Kia Las Vegas",
    "KTJ": "Kia St George", "ZHI": "Land Rover Henderson", "HUH": "Land Rover Reno", "IFA": "Lexus", "YH4": "Lincoln ",
    "FVQ": "Mazda Henderson", "Z7Q": "Motor Company", "SYT": "Nissan", "H8E": "Subaru Of Las Vegas",
    "XUW": "Subaru Prescott", "HP3": "Subaru St. George", "DAE": "Toyota Flagstaff", "Z6C": "Toyota Henderson",
    "XL1": "Toyota Prescott", "Q9I": "Toyota Spokane", "HR7": "Volkswagen Henderson", "Z06": "Volkswagen St. George",
    "4MZ": "Volvo Cars Las Vegas",
}

# ---------- load ----------
src = openpyxl.load_workbook(A.adp, data_only=True)
if "Sheet1" in src.sheetnames:
    for code, store in src["Sheet1"].iter_rows(min_row=1, max_col=2, values_only=True):
        if code and store and str(code).strip() != "Company Code":
            CODE_MAP[str(code).strip()] = store
ws1 = src["1"] if "1" in src.sheetnames else src.worksheets[0]
hdr = [c.value for c in ws1[1]]
ix = {h: hdr.index(h) for h in ("Company Code", "Payroll Name", "Job Title Description", "Gross Pay", "Position Status", "Hire/Rehire Date")}
rows, dropped, unmapped = [], 0, set()
for r in ws1.iter_rows(min_row=2, values_only=True):
    title = r[ix["Job Title Description"]]
    if not title:  # subtotal / footer rows
        continue
    code = str(r[ix["Company Code"]]).strip()
    if str(title).upper() in EXCL:
        dropped += 1; continue
    if code not in CODE_MAP:
        unmapped.add(code); continue
    hire = r[ix["Hire/Rehire Date"]]
    rows.append(dict(code=code, store=CODE_MAP[code], name=r[ix["Payroll Name"]], title=str(title),
                     status=r[ix["Position Status"]], hire=hire, gross=float(r[ix["Gross Pay"]] or 0)))
if unmapped:
    raise SystemExit(f"STOP: unmapped ADP company codes {sorted(unmapped)} - add them to CODE_MAP / Sheet1 before building")
rows.sort(key=lambda d: (d["code"], d["name"]))
for d in rows: d["orig_title"] = d["title"]
moved = []
if A.qlik:
    import json, os, sys
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from classify import classify, INTERNAL_TITLE
    QJ = json.load(open(A.qlik))
    if QJ.get("advisor_internal"):
        adp = [(d["code"], d["store"], d["name"], d["title"], i) for i, d in enumerate(rows)]
        for rec in classify(QJ["advisor_internal"]["rows"], adp, int(QJ["months"]), QJ.get("advisor_overrides")):
            if rec["affects"]:
                d = rows[rec["adp_rec"][4]]; d["title"] = INTERNAL_TITLE; moved.append(f'{d["store"].strip()}: {d["name"]}')
N = len(rows) + 1  # last Raw Data row

# ---------- styles ----------
NAVY, GOLD, LINK = "FF1A2744", "FFC9A04B", "FF1155CC"
thin = Side(style="thin", color="FFB7B7B7")  # light-gray gridlines, as in the original build
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF"); HDR_FILL = PatternFill("solid", fgColor=NAVY)
HDR_AL = Alignment(horizontal="center", vertical="center", wrap_text=True)
BASE = Font(name="Calibri", size=11); BOLD = Font(name="Calibri", size=11, bold=True)
GOLDF = PatternFill("solid", fgColor=GOLD)
CUR, CNT, PCT = "\\$#,##0.00", "#,##0", "0.0%"

def header(ws, row, heads, box=True):
    for i, h in enumerate(heads, 1):
        c = ws.cell(row, i, h); c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = HDR_AL
        if box: c.border = BOX

def put(ws, coord, v, fmt="General", font=BASE, fill=None, box=True):
    c = ws[coord]; c.value = v; c.number_format = fmt; c.font = font
    if fill: c.fill = fill
    if box: c.border = BOX
    return c

def widths(ws, w):
    for col, x in w.items(): ws.column_dimensions[col].width = x

def q(sheet):  # quoted sheet reference
    return f"'{sheet}'" if any(ch in sheet for ch in " &.-'") else sheet

RD = "'Raw Data'"
wb = openpyxl.Workbook(); wb.remove(wb.active)

# ---------- Raw Data ----------
rd = wb.create_sheet("Raw Data")
header(rd, 1, ["Company Code", "Store Name", "Payroll Name", "Job Title Description", "Position Status",
               "Hire/Rehire Date", "Gross Pay", "Months Employed (This Period)", "Included in Averages", "Average Monthly Pay",
               "Original Job Title (ADP)"], box=False)
for i, d in enumerate(rows, 2):
    for col, v, fmt in (("A", d["code"], "General"), ("B", d["store"], "General"), ("C", d["name"], "General"),
                        ("D", d["title"], "General"), ("E", d["status"], "General"), ("F", d["hire"], "mm-dd-yy"),
                        ("G", d["gross"], CUR),
                        ("H", f"=MAX(0,({DPE}-MAX(F{i},{DPS})+1))/30.4", "0.00"),
                        ("I", f'=IF(F{i}<={DPE},"Yes","No")', "General"),
                        ("J", f'=IF(AND(I{i}="Yes",H{i}>0),G{i}/H{i},"N/A")', CUR),
                        ("K", d["orig_title"], "General")):
        put(rd, f"{col}{i}", v, fmt, box=False)
rd.auto_filter.ref = f"A1:K{N}"; rd.freeze_panes = "A2"
widths(rd, dict(A=15, B=30, C=36, D=28, E=18, F=19, G=14, H=32, I=23, J=22, K=28))

# ---------- Group Position Summary ----------
def included(d):
    h = d["hire"]; h = h.date() if isinstance(h, dt.datetime) else h
    return isinstance(h, dt.date) and h <= PE
# positions / store x position combos are built from employees included in averages only
# (a combo whose only employees were hired after period end is left off, as in the delivered build)
titles = sorted({d["title"] for d in rows if included(d)})
gp = wb.create_sheet("Group Position Summary")
header(gp, 1, ["Job Title Description", "Employee Count", "Total Monthly Pay", "Avg Monthly Pay per Employee"], box=False)
for i, t in enumerate(titles, 2):
    put(gp, f"A{i}", t)
    put(gp, f"B{i}", f"=COUNTIFS({RD}!$D$2:$D${N},A{i},{RD}!$I$2:$I${N},\"Yes\")", CNT)
    put(gp, f"C{i}", f"=SUMIFS({RD}!$J$2:$J${N},{RD}!$D$2:$D${N},A{i},{RD}!$I$2:$I${N},\"Yes\")", CUR)
    put(gp, f"D{i}", f"=IF(B{i}=0,0,C{i}/B{i})", CUR)
GL = len(titles) + 1; GT = GL + 1
put(gp, f"A{GT}", "Grand Total", font=BOLD, fill=GOLDF)
put(gp, f"B{GT}", f"=SUBTOTAL(109,B2:B{GL})", CNT, BOLD, GOLDF)
put(gp, f"C{GT}", f"=SUBTOTAL(109,C2:C{GL})", CUR, BOLD, GOLDF)
put(gp, f"D{GT}", f"=IF(B{GT}=0,0,C{GT}/B{GT})", CUR, BOLD, GOLDF)
t = Table(displayName="GroupPositionSummary", ref=f"A1:D{GT}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True); gp.add_table(t)
gp.row_dimensions[1].height = 30; gp.freeze_panes = "A2"; widths(gp, dict(A=30, B=15, C=17, D=18))
GPR = f"'Group Position Summary'!$A$2:$A${GL}"; GPD = f"'Group Position Summary'!$D$2:$D${GL}"

# ---------- Store & Position Summary ----------
by_store = defaultdict(set)
for d in rows:
    if included(d): by_store[d["store"]].add(d["title"])
stores = sorted(by_store)
sp = wb.create_sheet("Store & Position Summary")
header(sp, 1, ["Store Name", "Job Title Description", "Employee Count", "Total Monthly Pay", "Avg Monthly Pay per Employee"], box=False)
i = 2
for s in stores:
    for tt in sorted(by_store[s]):
        put(sp, f"A{i}", s); put(sp, f"B{i}", tt)
        put(sp, f"C{i}", f"=COUNTIFS({RD}!$B$2:$B${N},A{i},{RD}!$D$2:$D${N},B{i},{RD}!$I$2:$I${N},\"Yes\")", CNT)
        put(sp, f"D{i}", f"=SUMIFS({RD}!$J$2:$J${N},{RD}!$B$2:$B${N},A{i},{RD}!$D$2:$D${N},B{i},{RD}!$I$2:$I${N},\"Yes\")", CUR)
        put(sp, f"E{i}", f"=IF(C{i}=0,0,D{i}/C{i})", CUR)
        i += 1
ST = i
put(sp, f"A{ST}", "Grand Total", font=BOLD, fill=GOLDF); put(sp, f"B{ST}", None, font=BOLD, fill=GOLDF)
put(sp, f"C{ST}", f"=SUBTOTAL(109,C2:C{ST-1})", CNT, BOLD, GOLDF)
put(sp, f"D{ST}", f"=SUBTOTAL(109,D2:D{ST-1})", CUR, BOLD, GOLDF)
put(sp, f"E{ST}", f"=IF(C{ST}=0,0,D{ST}/C{ST})", CUR, BOLD, GOLDF)
t = Table(displayName="StorePositionSummary", ref=f"A1:E{ST}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True); sp.add_table(t)
sp.row_dimensions[1].height = 30; sp.freeze_panes = "A2"; widths(sp, dict(A=28, B=26, C=15, D=17, E=18))

# ---------- store sheets ----------
def dxf(font_rgb, fill_rgb):  # CF fills need bgColor as well as fgColor or Excel shows no fill
    return DifferentialStyle(font=Font(color=font_rgb), fill=PatternFill(fill_type="solid", fgColor=fill_rgb, bgColor=fill_rgb))

ranges = {}
for s in stores:
    tab = s.strip()
    ws = wb.create_sheet(tab)
    c = ws["A1"]; c.value = "← Back to Master Summary"; c.hyperlink = Hyperlink(ref="A1", location="'Master Summary'!A1"); c.font = Font(name="Calibri", size=11, color=LINK)
    ws["A3"] = s; ws["A3"].font = Font(name="Calibri", size=14, bold=True, color=NAVY); ws.merge_cells("A3:H3")
    ws.row_dimensions[3].height = 18.75
    header(ws, 5, ["Job Title Description", "Employee Count", f"{tab} Avg Monthly Pay", "Group Avg Monthly Pay",
                   "Variance ($) per Employee", "Variance (%)", "Status", "Total Variance ($) - All Employees in Position"])
    ws.row_dimensions[5].height = 60
    sq = s.replace('"', '""')
    ts = sorted(by_store[s]); k = 5 + len(ts)
    for r, tt in enumerate(ts, 6):
        crit = f"{RD}!$B$2:$B${N},\"{sq}\",{RD}!$D$2:$D${N},\"{tt}\",{RD}!$I$2:$I${N},\"Yes\""
        put(ws, f"A{r}", tt)
        put(ws, f"B{r}", f"=COUNTIFS({crit})", CNT)
        put(ws, f"C{r}", f"=IFERROR(SUMIFS({RD}!$J$2:$J${N},{crit})/B{r},0)", CUR)
        put(ws, f"D{r}", f"=INDEX({GPD},MATCH(A{r},{GPR},0))", CUR)
        put(ws, f"E{r}", f"=C{r}-D{r}", CUR)
        put(ws, f"F{r}", f"=IF(D{r}=0,0,E{r}/D{r})", PCT)
        put(ws, f"G{r}", f'=IF(E{r}>0,"Above Average",IF(E{r}<0,"Below Average","At Average"))')
        put(ws, f"H{r}", f"=E{r}*B{r}", CUR)
    t1, ab, be, nt = k + 2, k + 3, k + 4, k + 5
    for col in "ABCDEFGH": put(ws, f"{col}{t1}", None, font=BOLD, fill=GOLDF)
    ws[f"A{t1}"] = "TOTAL / SUMMARY"; ws[f"B{t1}"] = f"=SUM(B6:B{k})"; ws[f"B{t1}"].number_format = CNT
    ws[f"H{t1}"] = f"=SUM(H6:H{k})"; ws[f"H{t1}"].number_format = CUR
    for row, label, f, font, fill in (
            (ab, "Total $ Paid Above Group Average (headcount-weighted)", f'=SUMIF(H6:H{k},">0")',
             Font(name="Calibri", size=11, color="FF9C0006"), PatternFill("solid", fgColor="FFFFC7CE")),
            (be, "Total $ Paid Below Group Average (headcount-weighted)", f'=SUMIF(H6:H{k},"<0")',
             Font(name="Calibri", size=11, color="FF006100"), PatternFill("solid", fgColor="FFC6EFCE")),
            (nt, "Net Variance $ (headcount-weighted)", f"=H{ab}+H{be}", BOLD, GOLDF)):
        for col in "ABCDH": put(ws, f"{col}{row}", None, font=font, fill=fill)
        ws[f"A{row}"] = label; ws[f"H{row}"] = f; ws[f"H{row}"].number_format = CUR
        ws.merge_cells(f"A{row}:D{row}")
    ws.conditional_formatting.add(f"E6:G{k}", Rule(type="expression", dxf=dxf("FF9C0006", "FFFFC7CE"), formula=["$E6>0"]))
    ws.conditional_formatting.add(f"E6:G{k}", Rule(type="expression", dxf=dxf("FF006100", "FFC6EFCE"), formula=["$E6<0"]))
    ws.freeze_panes = "A6"
    widths(ws, dict(A=28, B=10, C=15, D=13, E=12, F=10, G=22, H=16))
    ranges[s] = (tab, k)

# ---------- store-level % above (for Store Ranker order) ----------
def mpay(d):
    if not isinstance(d["hire"], (dt.date, dt.datetime)): return None
    h = d["hire"].date() if isinstance(d["hire"], dt.datetime) else d["hire"]
    if h > PE: return None
    m = max(0, (PE - max(h, PS)).days + 1) / 30.4
    return d["gross"] / m if m > 0 else None
grp = defaultdict(list); sto = defaultdict(list)
for d in rows:
    p = mpay(d)
    if p is not None: grp[d["title"]].append(p); sto[(d["store"], d["title"])].append(p)
pct = {}
for s in stores:
    ts = sorted(by_store[s]); above = 0
    for tt in ts:
        sv = sum(sto[(s, tt)]) / len(sto[(s, tt)]) if sto[(s, tt)] else 0
        gv = sum(grp[tt]) / len(grp[tt]) if grp[tt] else 0
        above += (sv - gv) > 0
    pct[s] = above / len(ts) if ts else 0

# ---------- Master Summary / Store Ranker ----------
NOTE_MS = ('Click a store name to jump to its position-level detail. "Above Avg" = store pays more than the group average for that role; '
           '"Below Avg" = pays less than the group average. Excludes Express Service Manager and Lead Sales Porter positions, and excludes '
           f'employees hired after {PE_TXT} from all averages (see Raw Data "Included in Averages" flag). See the Store Ranker tab for stores sorted by % above average.')
NOTE_SR = ('Stores are ordered highest to lowest by the share of positions where store pay exceeds the group average. Excludes Express Service '
           f'Manager and Lead Sales Porter positions, and excludes employees hired after {PE_TXT} from all averages. Ranking reflects the data '
           'as of this build; re-run after major data changes to refresh order.')

def summary(name, title, note, order):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=NAVY); ws.merge_cells("A1:G1")
    ws.row_dimensions[1].height = 18.75
    ws["A2"] = note; ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="FF555555"); ws.merge_cells("A2:G2")
    header(ws, 4, ["Store Name", "Total Positions", "Positions Above Group Avg", "Positions Below Group Avg",
                   "% of Positions Above Avg", "Store Weighted Avg Monthly Pay", "Group Weighted Avg Monthly Pay"])
    ws.row_dimensions[4].height = 45
    for r, s in enumerate(order, 5):
        tab, k = ranges[s]; t = q(tab)
        c = put(ws, f"A{r}", s, font=Font(name="Calibri", size=11, color=LINK)); c.hyperlink = Hyperlink(ref=f"A{r}", location=f"'{tab}'!A1")
        put(ws, f"B{r}", f"=COUNTA({t}!$A$6:$A${k})", CNT)
        put(ws, f"C{r}", f'=COUNTIF({t}!$G$6:$G${k},"Above*")', CNT)
        put(ws, f"D{r}", f'=COUNTIF({t}!$G$6:$G${k},"Below*")', CNT)
        put(ws, f"E{r}", f"=IF(B{r}=0,0,C{r}/B{r})", PCT)
        put(ws, f"F{r}", f"=IFERROR(SUMPRODUCT({t}!$B$6:$B${k},{t}!$C$6:$C${k})/SUM({t}!$B$6:$B${k}),0)", CUR)
        put(ws, f"G{r}", f"=IFERROR(SUMPRODUCT({t}!$B$6:$B${k},{t}!$D$6:$D${k})/SUM({t}!$B$6:$B${k}),0)", CUR)
    L = 4 + len(order); T = L + 1
    put(ws, f"A{T}", "Grand Total / Group", font=BOLD, fill=GOLDF)
    for col, f, fmt in (("B", f"=SUM(B5:B{L})", CNT), ("C", f"=SUM(C5:C{L})", CNT), ("D", f"=SUM(D5:D{L})", CNT),
                        ("E", f"=IF(B{T}=0,0,C{T}/B{T})", PCT),
                        ("F", f"=SUMPRODUCT('Group Position Summary'!$B$2:$B${GL},'Group Position Summary'!$D$2:$D${GL})/SUM('Group Position Summary'!$B$2:$B${GL})", CUR),
                        ("G", f"=F{T}", CUR)):
        put(ws, f"{col}{T}", f, fmt, BOLD, GOLDF)
    ws.conditional_formatting.add(f"C5:C{L}", ColorScaleRule(start_type="min", start_color="FFFFFFFF", end_type="max", end_color="FFF8696B"))
    ws.conditional_formatting.add(f"D5:D{L}", ColorScaleRule(start_type="min", start_color="FFFFFFFF", end_type="max", end_color="FFF8696B"))
    ws.conditional_formatting.add(f"E5:E{L}", ColorScaleRule(start_type="min", start_color="FF63BE7B", mid_type="percentile", mid_value=50,
                                                             mid_color="FFFFEB84", end_type="max", end_color="FFF8696B"))
    ws.freeze_panes = "A5"; widths(ws, dict(A=28, B=11, C=13, D=13, E=12, F=15, G=13))

summary("Master Summary", "Store Pay-vs-Group-Average Master Summary", NOTE_MS, stores)
summary("Store Ranker", "Store Ranker — Ranked by % of Positions Above Group Average", NOTE_SR,
        sorted(stores, key=lambda s: (-pct[s], s.strip())))

# ---------- tab order ----------
order = ["Master Summary", "Store Ranker", "Raw Data", "Store & Position Summary", "Group Position Summary"] + [ranges[s][0] for s in stores]
wb._sheets = [wb[n] for n in order]
wb.active = 0
wb.save(A.out)
print(f"saved {A.out}: {len(rows)} employees (dropped {dropped} excluded-position rows), {len(stores)} stores, "
      f"{len(titles)} positions, {ST-2} store x position rows, {len(moved)} moved to INTERNAL SERVICE ADVISOR. NEXT: add_gross_workload.py, then recalc.py")
