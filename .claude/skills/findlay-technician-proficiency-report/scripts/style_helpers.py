import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

ARIAL = "Arial"
NAVY = "1A2744"
GOLD = "C9A04B"
WHITE = "FFFFFF"
BLACK = "000000"
YELLOW_FILL = "FFFF00"

GREEN_FONT = "006100"
GREEN_FILL = "C6EFCE"
YELLOW_FONT = "9C6500"
YELLOW_FILL_CF = "FFEB9C"
RED_FONT = "9C0006"
RED_FILL = "FFC7CE"

from config import load_config

_cfg = load_config()
NEW_PERIOD_LABEL = _cfg["new_period_label"]
NEW_PERIOD_SUBTITLE = f"{_cfg['qlik_date_start']} - {_cfg['qlik_date_end']}"

CATEGORY_ORDER = ["Service Technician", "Express Technician", "Service Team Leader", "Shop Foreman",
                   "Body Shop Tech", "Body Shop Painter", "Body Shop Helper", "Detailer", "Tinter"]

# Categories included in Tech Ranker (and now, Current Store Proficiency % on each store sheet):
# actual producing technicians only -- excludes Body Shop, Detailer, Tinter, and unmapped/pooled codes.
PROF_CATEGORIES = {"Service Technician", "Express Technician", "Service Team Leader", "Shop Foreman"}

COL = {
    "A": 1, "B": 2, "C": 3, "D": 4, "E": 5, "F": 6, "G": 7, "H": 8, "I": 9,
    "J": 10, "K": 11, "L": 12, "M": 13, "N": 14, "O": 15, "P": 16, "Q": 17, "R": 18, "S": 19,
}

DATA_HEADERS = ["Tech No", "Name", "Position", "Sold Hours", "RO Count", "Avg Hrs/RO",
                "Actual Hours (Reg+OT)", "Proficiency %", "Status", "Labor Sale", "Labor Gross",
                "GP%", "ELR", None, "Current\nProficiency %", "Target\nProficiency %",
                "Proficiency\nIncrease", "Additional\nSold Hours", "Additional\nLabor Gross"]


def proficiency(row):
    """Sold Hours / Actual Hours, matching the workbook's own IFERROR(...,0) convention."""
    actual = row.get("actual_hours")
    return (row["sold_hours"] / actual) if actual else 0.0


def sort_key_proficiency(row):
    """Sort key for ordering named technicians within a category: highest proficiency first."""
    return -proficiency(row)


def qualifies_for_prof_subset(row, category):
    """Matches the Proficiency sheet's row-2 methodology: only Service Technician, Express
    Technician, Service Team Leader, and Shop Foreman count, and excludes outlier records --
    Service Team Leaders/Shop Foremen at or below 20% proficiency, and Service/Express
    Technicians at or below 2% proficiency."""
    if category not in PROF_CATEGORIES:
        return False
    prof = proficiency(row)
    if category in ("Service Team Leader", "Shop Foreman"):
        return prof > 0.20
    else:
        return prof > 0.02

def f(name=ARIAL, size=11, bold=False, color=None, underline=None, italic=False):
    kwargs = dict(name=name, size=size, bold=bold, italic=italic)
    if color:
        kwargs["color"] = color
    if underline:
        kwargs["underline"] = underline
    return Font(**kwargs)


def fill(color):
    return PatternFill("solid", fgColor=color)


def thin_border():
    side = Side(style="thin", color="CCCCCC")
    return Border(top=side, bottom=side, left=side, right=side)


def center():
    return Alignment(horizontal="center", vertical="center")


def wrap_center():
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def gen_left():
    return Alignment(horizontal="general")


def set_cell(ws, row, col, value, number_format=None, font=None, fill_color=None, align=None, border=None):
    cell = ws.cell(row=row, column=col, value=value)
    if number_format:
        cell.number_format = number_format
    if font:
        cell.font = font
    if fill_color:
        cell.fill = fill(fill_color)
    if align:
        cell.alignment = align
    if border:
        cell.border = border
    return cell


def add_status_cf(ws, col_letter, row_start, row_end, kind="pct"):
    """kind: 'pct' for numeric >=1/>=0.85 thresholds (O column style),
    'gp' for numeric >=0.8/>=0.75 thresholds (L column style),
    'status_text' for I column (Under/On/Above Target text match)."""
    rng = f"{col_letter}{row_start}:{col_letter}{row_end}"
    if kind == "status_text":
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'${col_letter}{row_start}="Under Target"'],
                              font=Font(color=RED_FONT), fill=fill(RED_FILL)))
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'${col_letter}{row_start}="On Target"'],
                              font=Font(color=YELLOW_FONT), fill=fill(YELLOW_FILL_CF)))
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'${col_letter}{row_start}="Above Target"'],
                              font=Font(color=GREEN_FONT), fill=fill(GREEN_FILL)))
    elif kind == "gp":
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'{col_letter}{row_start}>=0.8'],
                              font=Font(color=GREEN_FONT), fill=fill(GREEN_FILL)))
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'AND({col_letter}{row_start}>=0.75,{col_letter}{row_start}<0.8)'],
                              font=Font(color=YELLOW_FONT), fill=fill(YELLOW_FILL_CF)))
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'{col_letter}{row_start}<0.75'],
                              font=Font(color=RED_FONT), fill=fill(RED_FILL)))
    elif kind == "pct":
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'${col_letter}{row_start}>=1'],
                              font=Font(color=GREEN_FONT), fill=fill(GREEN_FILL)))
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'AND(${col_letter}{row_start}>=0.85,${col_letter}{row_start}<1)'],
                              font=Font(color=YELLOW_FONT), fill=fill(YELLOW_FILL_CF)))
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'${col_letter}{row_start}<0.85'],
                              font=Font(color=RED_FONT), fill=fill(RED_FILL)))


def add_posneg_cf(ws, col_letter, row_start, row_end, black_font=False):
    rng = f"{col_letter}{row_start}:{col_letter}{row_end}"
    pos_font = Font(color=BLACK, bold=True) if black_font else Font(color=GREEN_FONT)
    neg_font = Font(color=BLACK, bold=True) if black_font else Font(color=RED_FONT)
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'AND(ISNUMBER(${col_letter}{row_start}),${col_letter}{row_start}>0)'],
                          font=pos_font, fill=fill(GREEN_FILL)))
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'AND(ISNUMBER(${col_letter}{row_start}),${col_letter}{row_start}<0)'],
                          font=neg_font, fill=fill(RED_FILL)))


def add_whatif_flag_cf(ws, row_start, row_end):
    rng = f"B{row_start}:B{row_end}"
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'$P{row_start}<>""'], fill=fill(YELLOW_FILL)))


def add_category_block_cf(ws, row_start, row_end):
    add_status_cf(ws, "I", row_start, row_end, kind="status_text")
    add_status_cf(ws, "L", row_start, row_end, kind="gp")
    add_status_cf(ws, "O", row_start, row_end, kind="pct")
    add_posneg_cf(ws, "Q", row_start, row_end, black_font=False)
    add_posneg_cf(ws, "R", row_start, row_end, black_font=True)
    add_posneg_cf(ws, "S", row_start, row_end, black_font=False)
    add_whatif_flag_cf(ws, row_start, row_end)
